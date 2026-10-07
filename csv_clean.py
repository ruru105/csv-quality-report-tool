import argparse
import datetime
import re
import sys
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import IllegalCharacterError

EXCEL_SUFFIXES = (".xlsx", ".xlsm")
# 保存できるExcelは .xlsx だけ(.xlsm はマクロ付きの形式で、この形では作れない)。
UNSUPPORTED_OUTPUT_SUFFIXES = (".xls", ".xlsm")


DEFAULT_OUTPUT_FILE = "cleaned.csv"

# カンマ区切りとして読んだ結果が1列だけになったとき、
# 見出しにこれらの文字が含まれていれば、その文字を区切り文字として読み直す。
CANDIDATE_DELIMITERS = ["\t", ";", "|"]


def parse_args(argv):
    """コマンドで指定された、整形するCSVと出力先を読み取る。"""
    parser = argparse.ArgumentParser(
        description=(
            "CSVまたはExcel(.xlsx)ファイルを整形して、別のファイルとして保存します。"
            "保存先の名前が .xlsx で終わるときはExcel、それ以外はCSVで保存します。"
            "元のファイルは書き換えません。"
        )
    )
    parser.add_argument(
        "input_file",
        help="整形するCSVまたはExcel(.xlsx)ファイル",
    )
    parser.add_argument(
        "-o",
        "--output",
        default=DEFAULT_OUTPUT_FILE,
        help=(
            f"整形結果を保存するファイル(省略すると {DEFAULT_OUTPUT_FILE})。"
            ".xlsx で終わる名前を指定すると、Excelで保存する"
        ),
    )
    parser.add_argument(
        "--drop",
        help=(
            "削除する列名。複数指定するときはカンマで区切る"
            "(例: --drop 備考,メモ)"
        ),
    )
    parser.add_argument(
        "--where",
        help=(
            "残す行の条件を「列名=値」で指定する。値は完全に一致する行だけ残す"
            "(例: --where 部署=Care)"
        ),
    )
    parser.add_argument(
        "--sort",
        help=(
            "並べ替えに使う列名。数字だけの列は数の大きさで並べる"
            "(例: --sort 金額)。空欄の行は、いつも最後に置く"
        ),
    )
    parser.add_argument(
        "--desc",
        action="store_true",
        help="--sort を、大きい順(降順)にする",
    )
    parser.add_argument(
        "--rename",
        help=(
            "列名を変える。「元の名前=新しい名前」の形で、複数はカンマで区切る"
            "(例: --rename 名前=氏名,部署=所属)。--where・--drop・--sort には"
            "元の列名を書く"
        ),
    )
    parser.add_argument(
        "--sheet",
        help=(
            "Excelファイルを整形するとき、読み込むシートの名前"
            "(省略すると、いちばん左のシート)"
        ),
    )
    return parser.parse_args(argv)


def _excel_cell_to_text(value):
    """Excelのセルの値を、見た目に近い文字にする。空欄は空文字にする。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, datetime.date):
        return value.isoformat()
    return str(value)


def read_excel_as_text(file_path, sheet_name=None):
    """Excel(.xlsx)の1つのシートを、すべて文字のまま読み込む。

    1行目を見出しとして扱う。数式のセルは、Excelが保存した計算結果を読む。
    日付は「2026-10-01」の形、空欄は空欄にする。末尾の空の行・列は読み飛ばす。
    見出しが空の列や、同じ名前の見出しがあるときはエラーにする。
    """
    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
    except Exception as error:
        raise ValueError(
            "Excelファイルとして開けませんでした。"
            "パスワード付き・破損・形式違いのファイルは読めません。"
        ) from error

    try:
        if sheet_name is None:
            worksheet = workbook.worksheets[0]
        elif sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
        else:
            raise ValueError(
                f"シート「{sheet_name}」が見つかりません。"
                "あるシート：" + "、".join(workbook.sheetnames)
            )
        rows = [
            [_excel_cell_to_text(cell) for cell in row]
            for row in worksheet.iter_rows(values_only=True)
        ]
    finally:
        workbook.close()

    while rows and not any(text != "" for text in rows[-1]):
        rows.pop()
    if not rows:
        raise pd.errors.EmptyDataError("no data")

    width = max(
        (index + 1 for row in rows for index, text in enumerate(row) if text != ""),
        default=0,
    )
    rows = [(row + [""] * width)[:width] for row in rows]

    header = [name.strip() for name in rows[0]]
    if any(name == "" for name in header):
        raise ValueError("見出し(1行目)が空の列があります。見出しを書いてください。")
    if len(set(header)) != len(header):
        raise ValueError("同じ名前の見出しがあります。見出しを区別してください。")

    return pd.DataFrame(rows[1:], columns=header, dtype=str)


def read_input_as_text(file_path, sheet_name=None):
    """拡張子から、CSVかExcelかを判断して、文字のまま読み込む。"""
    suffix = Path(file_path).suffix.lower()
    if suffix == ".xls":
        raise ValueError(
            "古い形式のExcel(.xls)には対応していません。"
            "Excelで「.xlsx」として保存し直してください。"
        )
    if suffix in EXCEL_SUFFIXES:
        return read_excel_as_text(file_path, sheet_name)
    if sheet_name is not None:
        raise ValueError("--sheet はExcelファイルを読むときだけ指定できます。")
    return read_csv_as_text(file_path)


# 数として書き出してよいのは、Excelで数に直しても、元の文字に戻せる値だけ。
# 「007」「1.0」「090-1234」「12345678901234567890」などは、文字のままにする。
_PLAIN_NUMBER = re.compile(r"^-?(0|[1-9][0-9]*)(\.[0-9]*[1-9])?$")


def _excel_value(text):
    if _PLAIN_NUMBER.match(text) and len(text.replace("-", "").replace(".", "")) <= 15:
        number = float(text) if "." in text else int(text)
        if _excel_cell_to_text(number) == text:
            return number
    return text


def _visual_width(text):
    return sum(
        2 if unicodedata.east_asian_width(char) in ("F", "W") else 1
        for char in str(text)
    )


def write_excel(data, output_file):
    """整形したデータをExcel(.xlsx)で保存する。見出しは太字で固定し、絞り込みボタンを付ける。

    普通の数は数として、先頭が0の値や長すぎる数は文字として書き出す(値は変えない)。
    """
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "整形結果"
    worksheet.append(list(data.columns))
    for row in data.itertuples(index=False):
        worksheet.append([_excel_value(text) for text in row])
    # 「=」で始まる文字は、openpyxlが数式として保存してしまう。文字のまま保存するよう直す。
    for row in worksheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    worksheet.freeze_panes = "A2"
    if len(data.columns) > 0:
        worksheet.auto_filter.ref = worksheet.dimensions
    for column in worksheet.columns:
        longest = max(
            (_visual_width(cell.value) for cell in column if cell.value is not None),
            default=0,
        )
        worksheet.column_dimensions[column[0].column_letter].width = min(
            max(longest + 2, 8), 45
        )
    workbook.save(output_file)


def read_csv_as_text(file_path):
    """CSVを、すべて文字のまま読み込む。

    数字として読むと「007」が「7」に変わってしまうため、
    整形では見た目のまま扱う。空欄も空欄のまま残す。
    UTF-8とWindows用の文字コードに対応する。

    カンマ区切りとして読んだ結果、列が1つにまとまってしまったときは、
    見出しに含まれる記号(タブなど)から区切り文字を判断して読み直す。
    """
    for encoding in ("utf-8-sig", "cp932"):
        try:
            data = pd.read_csv(
                file_path,
                encoding=encoding,
                dtype=str,
                keep_default_na=False,
            )
        except UnicodeDecodeError:
            continue

        return _reread_if_wrong_delimiter(data, file_path, encoding)

    raise ValueError("CSVの文字コードを判定できませんでした。")


def _reread_if_wrong_delimiter(data, file_path, encoding):
    """1列だけの読み込み結果を確認し、区切り文字の判定違いなら読み直す。"""
    if len(data.columns) != 1:
        return data

    header = str(data.columns[0])
    for delimiter in CANDIDATE_DELIMITERS:
        if delimiter in header:
            return pd.read_csv(
                file_path,
                encoding=encoding,
                sep=delimiter,
                dtype=str,
                keep_default_na=False,
            )

    return data


def trim_spaces(data):
    """見出しとセルの前後の空白(全角の空白を含む)を消す。

    整形後のデータと、空白を消した箇所の数を返す。
    """
    trimmed = data.apply(lambda column: column.str.strip())
    changed_cells = int((trimmed != data).sum().sum())

    trimmed_names = [name.strip() for name in data.columns]
    changed_names = sum(
        1 for before, after in zip(data.columns, trimmed_names) if before != after
    )
    trimmed.columns = trimmed_names

    return trimmed, changed_cells + changed_names


def remove_duplicate_rows(data):
    """完全に同じ内容の行を、最初の1行だけ残して取り除く。

    整形後のデータと、取り除いた行数を返す。
    """
    unique_data = data.drop_duplicates()
    return unique_data, len(data) - len(unique_data)


def parse_where(condition):
    """「列名=値」の形式を、列名と値に分ける。"""
    if "=" not in condition:
        raise ValueError(
            f'--where の指定「{condition}」が正しくありません。'
            "「列名=値」の形式で指定してください。"
        )
    column_name, _, value = condition.partition("=")
    return column_name, value


def filter_rows(data, column_name, value):
    """指定した列が、指定した値と完全に一致する行だけ残す。

    列名がCSVにないときはエラーにする。整形後のデータを返す。
    """
    if column_name not in data.columns:
        raise ValueError(f"列名「{column_name}」がCSVに見つかりません。")
    return data[data[column_name] == value]


def drop_columns(data, column_names):
    """指定した列名を、CSVから削除する。

    指定した列名のうち、存在しないものがあればエラーにする。
    削除したあとに列が1つも残らないときもエラーにする。
    整形後のデータを返す。
    """
    missing = [name for name in column_names if name not in data.columns]
    if missing:
        raise ValueError(
            "次の列名がCSVに見つかりません：" + "、".join(missing)
        )

    dropped_data = data.drop(columns=column_names)
    if dropped_data.shape[1] == 0:
        raise ValueError("すべての列を削除することはできません。")

    return dropped_data


def sort_rows(data, column_name, descending=False):
    """指定した列の値で並べ替える。

    その列の空欄でない値がすべて数として読めるときは、数の大きさで並べる
    (「100」が「30」より先に来ないようにするため)。そうでないときは文字の順で並べる。
    空欄の行は、昇順でも降順でも最後に置く。同じ値の行は、元の順番を保つ。
    列名がCSVにないときはエラーにする。並べ替えるだけで、値そのものは変えない。
    """
    if column_name not in data.columns:
        raise ValueError(f"並べ替えに使う列名「{column_name}」がCSVに見つかりません。")

    column = data[column_name]
    filled = column != ""
    numbers = pd.to_numeric(column.where(filled), errors="coerce")
    use_numbers = bool(filled.any()) and not numbers[filled].isna().any()

    key = numbers if use_numbers else column.where(filled)
    # 行番号(0から数えた位置)で並べ替える。絞り込みのあとは、元の行番号が飛び飛びに
    # なっているため、行のラベルではなく位置で指定する。
    ranking = pd.DataFrame({"key": key.to_numpy(), "position": range(len(data))})
    ordered = ranking.sort_values(
        by=["key", "position"],
        ascending=[not descending, True],
        na_position="last",
        kind="stable",
    )
    return data.iloc[ordered["position"].tolist()]


def parse_rename(text):
    """「元の名前=新しい名前,…」を、(元の名前, 新しい名前)の一覧にする。"""
    pairs = []
    for part in text.split(","):
        part = part.strip()
        if not part:
            continue
        if "=" not in part:
            raise ValueError(
                f'--rename の指定「{part}」が正しくありません。'
                "「元の名前=新しい名前」の形で指定してください。"
            )
        old, _, new = part.partition("=")
        old, new = old.strip(), new.strip()
        if not old or not new:
            raise ValueError(
                f'--rename の指定「{part}」が正しくありません。'
                "元の名前と新しい名前の両方を書いてください。"
            )
        pairs.append((old, new))
    if not pairs:
        raise ValueError("--rename に、変更する列名が指定されていません。")
    return pairs


def rename_columns(data, pairs):
    """列名を変える。値と列の並びは変えない。

    CSVにない列名、同じ列を2回変える指定、変えたあとの名前が重なる指定はエラーにする。
    """
    olds = [old for old, _ in pairs]
    missing = [old for old in olds if old not in data.columns]
    if missing:
        raise ValueError(
            "名前を変える列がCSVに見つかりません：" + "、".join(missing)
        )
    if len(set(olds)) != len(olds):
        raise ValueError("同じ列の名前を2回変える指定はできません。")

    mapping = dict(pairs)
    new_names = [mapping.get(name, name) for name in data.columns]
    if len(set(new_names)) != len(new_names):
        raise ValueError(
            "名前を変えたあとの列名が、ほかの列と重なります。別の名前にしてください。"
        )
    renamed = data.copy()
    renamed.columns = new_names
    return renamed


def main(argv=None):
    args = parse_args(argv)
    input_file = Path(args.input_file)
    output_file = Path(args.output)

    if not input_file.exists():
        print(f"エラー：{input_file} が見つかりません。")
        return 1

    if output_file.suffix.lower() in UNSUPPORTED_OUTPUT_SUFFIXES:
        print(
            f"エラー：保存先の形式「{output_file.suffix}」には対応していません。"
            "Excelで保存するときは「.xlsx」、そのほかはCSVの名前にしてください。"
        )
        return 1

    if input_file.resolve() == output_file.resolve():
        print(
            "エラー：保存先が、整形するCSVと同じです。"
            "元のファイルを書き換えないように、別の名前を指定してください。"
        )
        return 1

    try:
        data = read_input_as_text(input_file, args.sheet)
    except pd.errors.EmptyDataError:
        print(f"エラー：{input_file} にデータがありません。")
        return 1
    except ValueError as error:
        print(f"エラー：{input_file} を読み込めませんでした。({error})")
        return 1

    row_count = len(data)

    data, space_count = trim_spaces(data)
    data, duplicate_count = remove_duplicate_rows(data)

    if args.where:
        try:
            column_name, value = parse_where(args.where)
            data = filter_rows(data, column_name, value)
        except ValueError as error:
            print(f"エラー：{error}")
            return 1

    dropped_columns = []
    if args.drop:
        dropped_columns = [
            name.strip() for name in args.drop.split(",") if name.strip()
        ]
        try:
            data = drop_columns(data, dropped_columns)
        except ValueError as error:
            print(f"エラー：{error}")
            return 1

    sort_note = ""
    if args.sort:
        try:
            data = sort_rows(data, args.sort.strip(), args.desc)
        except ValueError as error:
            print(f"エラー：{error}")
            return 1
        sort_note = f"{args.sort.strip()}（{'降順' if args.desc else '昇順'}）"
    elif args.desc:
        print("エラー：--desc は --sort と一緒に指定してください。")
        return 1

    renamed_pairs = []
    if args.rename:
        try:
            renamed_pairs = parse_rename(args.rename)
            data = rename_columns(data, renamed_pairs)
        except ValueError as error:
            print(f"エラー：{error}")
            return 1

    try:
        if output_file.suffix.lower() in EXCEL_SUFFIXES:
            write_excel(data, output_file)
        else:
            data.to_csv(output_file, index=False, encoding="utf-8-sig")
    except IllegalCharacterError:
        print(
            "エラー：Excelに書き込めない文字(制御文字)がデータに含まれています。"
            "元のデータを確認してください。"
        )
        return 1
    except OSError:
        print(
            f"エラー：{output_file} に保存できませんでした。"
            "Excelなどで開いていないか、保存先のフォルダがあるか確認してください。"
        )
        return 1

    print(f"完了：{output_file} を作成しました。")
    stats = (
        f"データ件数={row_count}→{len(data)} / "
        f"空白を消した箇所={space_count} / "
        f"取り除いた重複行={duplicate_count}"
    )
    if dropped_columns:
        stats += " / 削除した列=" + "、".join(dropped_columns)
    if sort_note:
        stats += " / 並べ替え=" + sort_note
    if renamed_pairs:
        stats += " / 列名の変更=" + "、".join(f"{o}→{n}" for o, n in renamed_pairs)
    print(stats)
    return 0


if __name__ == "__main__":
    sys.exit(main())
