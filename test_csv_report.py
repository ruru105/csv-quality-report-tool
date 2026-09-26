import pandas as pd
from openpyxl import load_workbook

from csv_report import read_csv_safely, main, visual_width


def test_read_csv_utf8(tmp_path):
    csv_file = tmp_path / "utf8.csv"
    csv_file.write_text(
        "id,name\n1,Aoki\n2,Sato\n",
        encoding="utf-8-sig",
    )

    data = read_csv_safely(csv_file)

    assert len(data) == 2
    assert list(data.columns) == ["id", "name"]


def test_read_csv_cp932(tmp_path):
    csv_file = tmp_path / "cp932.csv"
    csv_file.write_text(
        "id,name\n1,青木\n2,佐藤\n",
        encoding="cp932",
    )

    data = read_csv_safely(csv_file)

    assert len(data) == 2
    assert data.loc[0, "name"] == "青木"


def test_read_csv_tab_separated(tmp_path):
    # FX取引ソフトの出力など、タブ区切りのファイルも正しく列を分けて読めること。
    csv_file = tmp_path / "tab.csv"
    csv_file.write_text(
        "id\tname\n1\tAoki\n2\tSato\n",
        encoding="utf-8-sig",
    )

    data = read_csv_safely(csv_file)

    assert list(data.columns) == ["id", "name"]
    assert len(data) == 2


def test_read_csv_semicolon_separated(tmp_path):
    csv_file = tmp_path / "semi.csv"
    csv_file.write_text(
        "id;name\n1;Aoki\n2;Sato\n",
        encoding="utf-8-sig",
    )

    data = read_csv_safely(csv_file)

    assert list(data.columns) == ["id", "name"]
    assert len(data) == 2


def test_read_csv_genuinely_single_column_stays_single_column(tmp_path):
    # 区切り文字を含まない、もともと1列だけのCSVを誤って分割しないこと。
    csv_file = tmp_path / "single.csv"
    csv_file.write_text(
        "name\nAoki\nSato\n",
        encoding="utf-8-sig",
    )

    data = read_csv_safely(csv_file)

    assert list(data.columns) == ["name"]
    assert len(data) == 2


def test_report_generation(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    sample = pd.DataFrame(
        {
            "employee_id": [1001, 1002, 1003, 1002, 1004],
            "name": ["Aoki", "Sato", "Tanaka", "Sato", "Suzuki"],
            "department": ["Care", "Office", None, "Office", "Care"],
            "hours_worked": [8, 7.5, 8, 7.5, None],
        }
    )
    sample.to_csv("sample.csv", index=False, encoding="utf-8-sig")

    result = main()

    assert result == 0

    report_file = tmp_path / "report.xlsx"
    assert report_file.exists()

    workbook = load_workbook(report_file, data_only=True)

    assert workbook.sheetnames == [
        "検査結果",
        "列情報",
        "欠損行",
        "重複行",
        "元データ",
    ]

    sheet = workbook["検査結果"]

    assert sheet["A2"].value == "検査したCSV"
    assert sheet["B2"].value == "sample.csv"
    assert sheet["A3"].value == "検査日時"
    assert sheet["B3"].value  # 日時が何かしら入っていること(値そのものは時刻依存のため確認しない)
    assert sheet["B4"].value == 5
    assert sheet["B5"].value == 4
    assert sheet["B6"].value == 2
    assert sheet["B7"].value == 1
    assert sheet["B8"].value == "問題あり"


def test_report_with_specified_files(tmp_path):
    # sample.csv ではない名前のCSVを、コマンドで指定して検査できること。
    csv_file = tmp_path / "my_data.csv"
    csv_file.write_text(
        "id,name\n1,Aoki\n2,\n2,\n",
        encoding="utf-8-sig",
    )
    output_file = tmp_path / "my_report.xlsx"

    result = main([str(csv_file), "-o", str(output_file)])

    assert result == 0
    assert output_file.exists()

    sheet = load_workbook(output_file, data_only=True)["検査結果"]

    assert sheet["B2"].value == str(csv_file)
    assert sheet["B4"].value == 3
    assert sheet["B5"].value == 2
    assert sheet["B6"].value == 2
    assert sheet["B7"].value == 1
    assert sheet["B8"].value == "問題あり"


def test_missing_input_file_shows_error(tmp_path, capsys):
    # 存在しないCSVを指定したとき、分かりやすいエラーを出してレポートを作らないこと。
    missing_file = tmp_path / "no_such_file.csv"
    output_file = tmp_path / "report.xlsx"

    result = main([str(missing_file), "-o", str(output_file)])

    assert result == 1
    assert "見つかりません" in capsys.readouterr().out
    assert not output_file.exists()


def test_empty_csv_shows_error(tmp_path, capsys):
    # 空のCSVを指定したとき、分かりやすいエラーを出してレポートを作らないこと。
    empty_file = tmp_path / "empty.csv"
    empty_file.write_text("", encoding="utf-8-sig")
    output_file = tmp_path / "report.xlsx"

    result = main([str(empty_file), "-o", str(output_file)])

    assert result == 1
    assert "データがありません" in capsys.readouterr().out
    assert not output_file.exists()


def test_broken_csv_shows_error(tmp_path, capsys):
    # 列数が行によって違う、壊れた形式のCSVを指定したとき、
    # 分かりやすいエラーを出してレポートを作らないこと。
    broken_file = tmp_path / "broken.csv"
    broken_file.write_text(
        "id,name\n1,Aoki\n2,Sato,Extra\n",
        encoding="utf-8-sig",
    )
    output_file = tmp_path / "report.xlsx"

    result = main([str(broken_file), "-o", str(output_file)])

    assert result == 1
    assert "読み込めませんでした" in capsys.readouterr().out
    assert not output_file.exists()


def test_visual_width_counts_full_width_characters_as_two():
    # 半角英字は1、日本語(全角)は2として数えること。
    assert visual_width("id") == 2
    assert visual_width("有効データ数") == 12


def test_status_cell_is_highlighted_when_problem_found(tmp_path):
    csv_file = tmp_path / "data.csv"
    csv_file.write_text(
        "id,name\n1,Aoki\n2,\n",
        encoding="utf-8-sig",
    )
    output_file = tmp_path / "report.xlsx"

    main([str(csv_file), "-o", str(output_file)])

    sheet = load_workbook(output_file)["検査結果"]
    status_cell = sheet["B8"]

    assert status_cell.value == "問題あり"
    assert status_cell.fill.fgColor.rgb == "00FFC7CE"
    assert status_cell.font.color.rgb == "009C0006"


def test_status_cell_is_highlighted_when_no_problem(tmp_path):
    csv_file = tmp_path / "data.csv"
    csv_file.write_text(
        "id,name\n1,Aoki\n2,Sato\n",
        encoding="utf-8-sig",
    )
    output_file = tmp_path / "report.xlsx"

    main([str(csv_file), "-o", str(output_file)])

    sheet = load_workbook(output_file)["検査結果"]
    status_cell = sheet["B8"]

    assert status_cell.value == "問題なし"
    assert status_cell.fill.fgColor.rgb == "00C6EFCE"
    assert status_cell.font.color.rgb == "00006100"
