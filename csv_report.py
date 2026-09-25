import argparse
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill


DEFAULT_INPUT_FILE = "sample.csv"
DEFAULT_OUTPUT_FILE = "report.xlsx"

# 総合判定のセルを目立たせる色(薄い塗りつぶし・濃い文字色)。
STATUS_STYLES = {
    "問題あり": {"fill": "FFC7CE", "font": "9C0006"},
    "問題なし": {"fill": "C6EFCE", "font": "006100"},
}


def parse_args(argv):
    """コマンドで指定された、検査するCSVと出力先を読み取る。"""
    parser = argparse.ArgumentParser(
        description="CSVファイルを検査して、結果をExcelレポートにまとめます。"
    )
    parser.add_argument(
        "input_file",
        nargs="?",
        default=DEFAULT_INPUT_FILE,
        help=f"検査するCSVファイル(省略すると {DEFAULT_INPUT_FILE})",
    )
    parser.add_argument(
        "-o",
        "--output",
        default=DEFAULT_OUTPUT_FILE,
        help=f"作成するExcelファイル(省略すると {DEFAULT_OUTPUT_FILE})",
    )
    return parser.parse_args(argv)


def read_csv_safely(file_path):
    """UTF-8とWindows用の文字コードに対応してCSVを読み込む。"""
    for encoding in ("utf-8-sig", "cp932"):
        try:
            return pd.read_csv(file_path, encoding=encoding)
        except UnicodeDecodeError:
            continue

    raise ValueError("CSVの文字コードを判定できませんでした。")


def visual_width(text):
    """全角文字を2、半角文字を1として数えた、見た目の幅を返す。

    日本語の見出しは、文字数だけで列幅を決めると右端が切れることがあるため、
    見た目の幅で列幅を決める。
    """
    width = 0
    for char in text:
        width += 2 if unicodedata.east_asian_width(char) in ("F", "W") else 1
    return width


def format_workbook(writer):
    """Excelレポートを読みやすく整える。"""
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for worksheet in writer.book.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for column_cells in worksheet.columns:
            max_width = max(
                visual_width(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells
            )
            column_letter = column_cells[0].column_letter
            worksheet.column_dimensions[column_letter].width = min(
                max(max_width + 2, 12), 45
            )


def highlight_summary_sheet(writer):
    """「検査結果」シートの総合判定を目立たせ、結果列の寄せをそろえる。"""
    worksheet = writer.book["検査結果"]

    for row in worksheet.iter_rows(min_row=2):
        label_cell, value_cell = row[0], row[1]
        # 数字と文字が混ざっていても、結果列の寄せをそろえる。
        value_cell.alignment = Alignment(horizontal="left")

        if label_cell.value == "総合判定":
            style = STATUS_STYLES.get(value_cell.value)
            if style:
                value_cell.fill = PatternFill("solid", fgColor=style["fill"])
                value_cell.font = Font(color=style["font"], bold=True)


def main(argv=None):
    # argv を渡さない(テストなど)ときは、指定なしとして動く。
    args = parse_args([] if argv is None else argv)
    input_file = Path(args.input_file)
    output_file = Path(args.output)

    if not input_file.exists():
        print(f"エラー：{input_file} が見つかりません。")
        return 1

    try:
        data = read_csv_safely(input_file)
    except pd.errors.EmptyDataError:
        print(f"エラー：{input_file} にデータがありません。")
        return 1
    except (pd.errors.ParserError, ValueError) as error:
        print(f"エラー：{input_file} を読み込めませんでした。({str(error).strip()})")
        return 1

    row_count = len(data)
    column_count = len(data.columns)
    missing_count = int(data.isna().sum().sum())
    duplicate_count = int(data.duplicated().sum())

    status = (
        "問題あり"
        if missing_count > 0 or duplicate_count > 0
        else "問題なし"
    )

    checked_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    summary = pd.DataFrame(
        {
            "確認項目": [
                "検査したCSV",
                "検査日時",
                "データ件数",
                "列数",
                "欠損セル数",
                "重複件数",
                "総合判定",
            ],
            "結果": [
                str(input_file),
                checked_at,
                row_count,
                column_count,
                missing_count,
                duplicate_count,
                status,
            ],
        }
    )

    column_info = pd.DataFrame(
        {
            "列名": data.columns,
            "データ型": [str(data[column].dtype) for column in data.columns],
            "有効データ数": [int(data[column].notna().sum()) for column in data.columns],
            "欠損数": [int(data[column].isna().sum()) for column in data.columns],
            "ユニーク数": [int(data[column].nunique(dropna=True)) for column in data.columns],
        }
    )

    missing_rows = data[data.isna().any(axis=1)]
    duplicate_rows = data[data.duplicated(keep=False)]

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="検査結果", index=False)
        column_info.to_excel(writer, sheet_name="列情報", index=False)
        missing_rows.to_excel(writer, sheet_name="欠損行", index=False)
        duplicate_rows.to_excel(writer, sheet_name="重複行", index=False)
        data.to_excel(writer, sheet_name="元データ", index=False)

        format_workbook(writer)
        highlight_summary_sheet(writer)

    print(f"完了：{output_file} を作成しました。")
    print(
        f"データ件数={row_count} / "
        f"欠損セル数={missing_count} / "
        f"重複件数={duplicate_count} / "
        f"判定={status}"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
