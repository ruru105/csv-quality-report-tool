"""検査結果をSQLite(標準ライブラリのsqlite3)に記録し、過去の結果と見比べる。"""
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

TABLE_SQL = """
CREATE TABLE IF NOT EXISTS inspections (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    checked_at  TEXT    NOT NULL,
    csv_path    TEXT    NOT NULL,
    file_name   TEXT    NOT NULL,
    status      TEXT    NOT NULL,
    rows        INTEGER,
    columns     INTEGER,
    missing     INTEGER,
    duplicates  INTEGER,
    violations  INTEGER,
    config_path TEXT,
    error       TEXT
)
"""

# 比較は、CSVの場所(フルパス)ごとに行う。別のフォルダの同名CSVは別物として扱う。
# 比較に使う数字の項目(表示名つき)。
COMPARE_FIELDS = [
    ("rows", "データ件数"),
    ("columns", "列数"),
    ("missing", "欠損セル数"),
    ("duplicates", "重複件数"),
    ("violations", "ルール違反数"),
]


class HistoryError(Exception):
    """記録用データベースを開けない・読めないときのエラー。"""


def open_history(db_path):
    """記録用データベースを開く(なければ作る)。開けなければ HistoryError。"""
    path = Path(db_path)
    try:
        if path.parent and not path.parent.exists():
            raise HistoryError(f"記録先のフォルダ {path.parent} が見つかりません。")
        connection = sqlite3.connect(str(path))
        connection.execute(TABLE_SQL)
        connection.commit()
    except sqlite3.Error as error:
        raise HistoryError(f"記録用データベース {path} を開けませんでした。({error})") from error
    return connection


def record_inspection(connection, csv_file, result, config_path=None):
    """1回の検査結果を1行として記録する。エラーだったCSVも「エラー」として残す。"""
    if result.get("error"):
        values = (result.get("checked_at") or _now(), str(Path(csv_file).resolve()), Path(csv_file).name,
                  "エラー", None, None, None, None, None, config_path, result["error"])
    else:
        values = (result["checked_at"], str(Path(csv_file).resolve()), Path(csv_file).name,
                  result["status"], result["rows"], result["columns"], result["missing"],
                  result["duplicates"], result["violations"], config_path, None)
    try:
        connection.execute(
            "INSERT INTO inspections (checked_at, csv_path, file_name, status, rows, columns,"
            " missing, duplicates, violations, config_path, error)"
            " VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            values,
        )
        connection.commit()
    except sqlite3.Error as error:
        raise HistoryError(f"検査結果を記録できませんでした。({error})") from error


def _now():
    from datetime import datetime
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def load_history(db_path, file_name=None, limit=None):
    """記録を古い順に返す(辞書のリスト)。file_name を指定すると、そのCSV名だけ。"""
    path = Path(db_path)
    if not path.is_file():
        raise HistoryError(f"記録用データベース {path} が見つかりません。")
    connection = open_history(path)
    try:
        connection.row_factory = sqlite3.Row
        sql = "SELECT * FROM inspections"
        params = []
        if file_name:
            sql += " WHERE file_name = ?"
            params.append(file_name)
        sql += " ORDER BY id"
        records = [dict(row) for row in connection.execute(sql, params)]
    except sqlite3.Error as error:
        raise HistoryError(f"記録を読み込めませんでした。({error})") from error
    finally:
        connection.close()
    # 前回との違いは、絞り込む前の記録全体で計算する(件数を絞っても、先頭が「初回」にならないように)
    last_by_file = {}
    for record in records:
        record["change"] = describe_changes(last_by_file.get(record["csv_path"]), record)
        last_by_file[record["csv_path"]] = record
    if limit is not None and limit > 0:
        # CSVごとに、新しい順で limit 件だけ残す(古い順の並びは保つ)。
        keep = set()
        seen = {}
        for record in reversed(records):
            seen[record["csv_path"]] = seen.get(record["csv_path"], 0) + 1
            if seen[record["csv_path"]] <= limit:
                keep.add(record["id"])
        records = [r for r in records if r["id"] in keep]
    return records


def describe_changes(previous, current):
    """同じCSVの、1つ前の記録との違いを文にして返す。違いがなければ「前回と同じ」。"""
    if previous is None:
        return "初回"
    parts = []
    if previous["status"] != current["status"]:
        parts.append(f"判定 {previous['status']}→{current['status']}")
    for key, label in COMPARE_FIELDS:
        before, after = previous[key], current[key]
        if before is None or after is None or before == after:
            continue
        sign = "+" if after > before else ""
        parts.append(f"{label} {before}→{after}({sign}{after - before})")
    return "、".join(parts) if parts else "前回と同じ"


def format_history(records):
    """記録を、画面に表示する行のリストにする(各行に、前回との違いを付ける)。"""
    lines = []
    last_by_file = {}
    for record in records:
        change = record.get("change") or describe_changes(last_by_file.get(record["csv_path"]), record)
        if record["status"] == "エラー":
            detail = f"エラー：{record['error']}"
        else:
            detail = (
                f"{record['status']} データ件数={record['rows']} 欠損セル数={record['missing']} "
                f"重複件数={record['duplicates']}"
            )
            if record["violations"] is not None:
                detail += f" ルール違反数={record['violations']}"
        lines.append(f"{record['checked_at']}  {record['file_name']}  {detail}  [{change}]")
        last_by_file[record["csv_path"]] = record
    return lines


EXPORT_SHEET_NAME = "履歴"
EXPORT_HEADERS = [
    "検査日時", "ファイル名", "判定", "データ件数", "列数", "欠損セル数",
    "重複件数", "ルール違反数", "前回との違い", "エラーの内容", "CSVの場所",
]


def export_history_excel(records, output_path):
    """記録をExcelに書き出す。1つのCSVの記録が2回分以上あれば、欠損・重複の推移グラフも付ける。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = EXPORT_SHEET_NAME
    sheet.append(EXPORT_HEADERS)

    last_by_file = {}
    for record in records:
        change = record.get("change") or describe_changes(last_by_file.get(record["csv_path"]), record)
        last_by_file[record["csv_path"]] = record
        sheet.append([
            record["checked_at"], record["file_name"], record["status"], record["rows"],
            record["columns"], record["missing"], record["duplicates"], record["violations"],
            change, record["error"] or "", record["csv_path"],
        ])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    widths = [20, 24, 10, 12, 8, 12, 10, 14, 50, 30, 50]
    for index, width in enumerate(widths):
        sheet.column_dimensions[chr(ord("A") + index)].width = width

    paths = {record["csv_path"] for record in records}
    numeric_rows = [r for r in records if r["status"] != "エラー"]
    if len(paths) == 1 and len(numeric_rows) >= 2:
        _add_trend_chart(workbook, numeric_rows)

    try:
        workbook.save(str(output_path))
    except OSError as error:
        raise HistoryError(f"Excelを保存できませんでした。({error})") from error


def _add_trend_chart(workbook, numeric_rows):
    """欠損セル数と重複件数が、検査のたびにどう変わったかを折れ線グラフにする。"""
    sheet = workbook.create_sheet("推移")
    sheet.append(["検査日時", "欠損セル数", "重複件数"])
    for record in numeric_rows:
        sheet.append([record["checked_at"], record["missing"], record["duplicates"]])
    chart = LineChart()
    chart.title = f"{numeric_rows[0]['file_name']} の推移"
    chart.y_axis.title = "件数"
    chart.x_axis.title = "検査日時"
    data = Reference(sheet, min_col=2, max_col=3, min_row=1, max_row=len(numeric_rows) + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=len(numeric_rows) + 1))
    chart.width, chart.height = 20, 9
    sheet.add_chart(chart, "E2")
    sheet.column_dimensions["A"].width = 20
