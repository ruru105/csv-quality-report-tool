import json

import pytest
from openpyxl import load_workbook

from csv_report import main, unique_report_name


def make_folder(tmp_path):
    """OK・重複あり・空・壊れ・大文字拡張子・CSV以外・サブフォルダ、を含むフォルダ。"""
    folder = tmp_path / "in"
    folder.mkdir()
    (folder / "ok.csv").write_text("id,name\n1,A\n2,B\n", encoding="utf-8")
    (folder / "dup.csv").write_text("id,name\n1,A\n1,A\n", encoding="utf-8")
    (folder / "empty.csv").write_text("", encoding="utf-8")
    (folder / "broken.csv").write_text("id,name\n1,A,x\n2\n", encoding="utf-8")
    (folder / "UPPER.CSV").write_text("x\n1\n", encoding="utf-8")
    (folder / "readme.txt").write_text("not a csv", encoding="utf-8")
    sub = folder / "sub"
    sub.mkdir()
    (sub / "inner.csv").write_text("id\n1\n", encoding="utf-8")
    return folder


def sheet_rows(path, sheet):
    return [[cell.value for cell in row] for row in load_workbook(path)[sheet].iter_rows()]


def summary_rows(out):
    rows = sheet_rows(out / "summary.xlsx", "一覧")
    header = rows[0]
    return {row[0]: dict(zip(header, row)) for row in rows[1:]}


def test_folder_creates_report_per_csv_and_summary(tmp_path):
    folder = make_folder(tmp_path)
    out = tmp_path / "out"
    main(["--folder", str(folder), "-o", str(out)])

    names = sorted(path.name for path in out.iterdir())
    assert names == ["UPPER_report.xlsx", "dup_report.xlsx", "ok_report.xlsx", "summary.xlsx"]


def test_folder_ignores_non_csv_and_subfolders_and_includes_upper_case(tmp_path):
    folder = make_folder(tmp_path)
    out = tmp_path / "out"
    main(["--folder", str(folder), "-o", str(out)])

    rows = summary_rows(out)
    assert set(rows) == {"broken.csv", "dup.csv", "empty.csv", "ok.csv", "UPPER.CSV"}


def test_summary_has_status_and_counts_per_file(tmp_path):
    folder = make_folder(tmp_path)
    out = tmp_path / "out"
    main(["--folder", str(folder), "-o", str(out)])

    rows = summary_rows(out)
    assert rows["ok.csv"]["判定"] == "問題なし"
    assert rows["ok.csv"]["データ件数"] == 2
    assert rows["dup.csv"]["判定"] == "問題あり"
    assert rows["dup.csv"]["重複件数"] == 1
    assert rows["dup.csv"]["レポート"] == "dup_report.xlsx"


def test_error_file_does_not_stop_others_and_is_listed(tmp_path):
    folder = make_folder(tmp_path)
    out = tmp_path / "out"
    code = main(["--folder", str(folder), "-o", str(out)])

    assert code == 1  # 検査できなかったCSVがあるので、終了コードは1
    rows = summary_rows(out)
    assert rows["empty.csv"]["判定"] == "エラー"
    assert "データがありません" in rows["empty.csv"]["内容"]
    assert rows["broken.csv"]["判定"] == "エラー"
    assert not (out / "empty_report.xlsx").exists()
    # エラーの後ろのCSVも検査されている
    assert rows["ok.csv"]["判定"] == "問題なし"
    assert (out / "ok_report.xlsx").exists()


def test_overall_sheet_counts(tmp_path):
    folder = make_folder(tmp_path)
    out = tmp_path / "out"
    main(["--folder", str(folder), "-o", str(out)])

    overall = dict(sheet_rows(out / "summary.xlsx", "全体結果")[1:])
    assert overall["CSVの数"] == 5
    assert overall["問題なし"] == 2
    assert overall["問題あり"] == 1
    assert overall["エラー(検査できなかった)"] == 2
    assert load_workbook(out / "summary.xlsx").sheetnames == ["全体結果", "一覧"]


def test_all_good_folder_returns_zero_even_with_problems(tmp_path):
    folder = tmp_path / "in"
    folder.mkdir()
    (folder / "dup.csv").write_text("id\n1\n1\n", encoding="utf-8")
    assert main(["--folder", str(folder), "-o", str(tmp_path / "out")]) == 0


def test_default_output_folder_is_reports(tmp_path, monkeypatch):
    folder = tmp_path / "in"
    folder.mkdir()
    (folder / "a.csv").write_text("id\n1\n", encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    assert main(["--folder", "in"]) == 0
    assert (tmp_path / "reports" / "summary.xlsx").exists()


def test_output_folder_is_created_including_parents(tmp_path):
    folder = tmp_path / "in"
    folder.mkdir()
    (folder / "a.csv").write_text("id\n1\n", encoding="utf-8")
    out = tmp_path / "x" / "y"
    assert main(["--folder", str(folder), "-o", str(out)]) == 0
    assert (out / "a_report.xlsx").exists()


def test_output_into_input_folder_does_not_touch_csv(tmp_path):
    folder = tmp_path / "in"
    folder.mkdir()
    csv_file = folder / "a.csv"
    csv_file.write_text("id\n1\n", encoding="utf-8")
    assert main(["--folder", str(folder), "-o", str(folder)]) == 0
    assert csv_file.read_text(encoding="utf-8") == "id\n1\n"


def test_missing_folder_is_an_error(tmp_path, capsys):
    assert main(["--folder", str(tmp_path / "none"), "-o", str(tmp_path / "out")]) == 1
    assert "見つかりません" in capsys.readouterr().out
    assert not (tmp_path / "out").exists()


def test_folder_without_csv_is_an_error(tmp_path, capsys):
    folder = tmp_path / "in"
    folder.mkdir()
    (folder / "memo.txt").write_text("x", encoding="utf-8")
    assert main(["--folder", str(folder), "-o", str(tmp_path / "out")]) == 1
    assert ".csv ファイルがありません" in capsys.readouterr().out
    assert not (tmp_path / "out").exists()


def test_folder_and_csv_file_together_is_an_error(tmp_path, capsys):
    folder = make_folder(tmp_path)
    assert main(["sample.csv", "--folder", str(folder), "-o", str(tmp_path / "out")]) == 1
    assert "同時に使えません" in capsys.readouterr().out


def test_folder_with_config_adds_rule_columns_and_sheets(tmp_path):
    folder = tmp_path / "in"
    folder.mkdir()
    (folder / "a.csv").write_text("id,name\n1,A\n1,B\n", encoding="utf-8")
    (folder / "b.csv").write_text("id,name\n1,A\n2,B\n", encoding="utf-8")
    rules = tmp_path / "rules.json"
    rules.write_text(json.dumps({"unique_columns": ["id"]}), encoding="utf-8")
    out = tmp_path / "out"
    main(["--folder", str(folder), "-c", str(rules), "-o", str(out)])

    rows = summary_rows(out)
    assert rows["a.csv"]["判定"] == "問題あり"
    assert rows["a.csv"]["ルール違反数"] == 1
    assert rows["b.csv"]["判定"] == "問題なし"
    assert "ルール判定" in load_workbook(out / "a_report.xlsx").sheetnames
    overall = dict(sheet_rows(out / "summary.xlsx", "全体結果")[1:])
    assert overall["検査ルール(設定ファイル)"] == str(rules)


def test_folder_without_config_has_no_rule_column(tmp_path):
    folder = make_folder(tmp_path)
    out = tmp_path / "out"
    main(["--folder", str(folder), "-o", str(out)])
    assert "ルール違反数" not in sheet_rows(out / "summary.xlsx", "一覧")[0]


def test_invalid_config_stops_before_creating_anything(tmp_path):
    folder = make_folder(tmp_path)
    rules = tmp_path / "bad.json"
    rules.write_text('{"unique_column": ["id"]}', encoding="utf-8")
    out = tmp_path / "out"
    assert main(["--folder", str(folder), "-c", str(rules), "-o", str(out)]) == 1
    assert not out.exists()


def test_unique_report_name_avoids_duplicates(tmp_path):
    used = set()
    first = unique_report_name(tmp_path / "a.csv", used)
    second = unique_report_name(tmp_path / "A.CSV", used)
    assert first == "a_report.xlsx"
    assert second == "A_report_2.xlsx"


# ---- 1ファイル検査で見つかった不具合の回帰テスト ----


def test_row_longer_than_header_is_an_error_not_a_crash(tmp_path, capsys):
    # 最初のデータ行が見出しより列が多いと、以前は Traceback で止まっていた。
    csv_file = tmp_path / "broken.csv"
    csv_file.write_text("id,name\n1,A,x\n2\n", encoding="utf-8")
    output = tmp_path / "out.xlsx"
    assert main([str(csv_file), "-o", str(output)]) == 1
    assert "見出しより多くの列" in capsys.readouterr().out
    assert not output.exists()
