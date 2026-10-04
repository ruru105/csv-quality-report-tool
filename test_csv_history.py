import sqlite3

from csv_history import (
    HistoryError,
    describe_changes,
    format_history,
    load_history,
    open_history,
)
from csv_report import main

import pytest


def write(path, text):
    path.write_text(text, encoding="utf-8")
    return path


def db_rows(db):
    connection = sqlite3.connect(str(db))
    connection.row_factory = sqlite3.Row
    try:
        return [dict(r) for r in connection.execute("SELECT * FROM inspections ORDER BY id")]
    finally:
        connection.close()


def test_single_check_is_recorded(tmp_path):
    csv = write(tmp_path / "a.csv", "id,name\n1,A\n1,A\n2,\n")
    db = tmp_path / "h.db"
    code = main([str(csv), "-o", str(tmp_path / "r.xlsx"), "--history", str(db)])
    assert code == 0
    rows = db_rows(db)
    assert len(rows) == 1
    assert rows[0]["file_name"] == "a.csv"
    assert rows[0]["status"] == "問題あり"
    assert (rows[0]["rows"], rows[0]["columns"], rows[0]["missing"], rows[0]["duplicates"]) == (3, 2, 1, 1)
    assert rows[0]["violations"] is None
    assert rows[0]["error"] is None


def test_each_run_appends_a_row(tmp_path):
    csv = write(tmp_path / "a.csv", "id\n1\n2\n")
    db = tmp_path / "h.db"
    for _ in range(3):
        main([str(csv), "-o", str(tmp_path / "r.xlsx"), "--history", str(db)])
    assert len(db_rows(db)) == 3


def test_rules_violations_are_recorded(tmp_path):
    csv = write(tmp_path / "a.csv", "id,name\n1,A\n2,\n")
    rules = write(tmp_path / "rules.json", '{"not_null_columns": ["name"]}')
    db = tmp_path / "h.db"
    main([str(csv), "-o", str(tmp_path / "r.xlsx"), "-c", str(rules), "--history", str(db)])
    row = db_rows(db)[0]
    assert row["violations"] == 1
    assert row["config_path"] == str(rules)


def test_error_csv_is_recorded_as_error(tmp_path):
    csv = write(tmp_path / "empty.csv", "")
    db = tmp_path / "h.db"
    code = main([str(csv), "-o", str(tmp_path / "r.xlsx"), "--history", str(db)])
    assert code == 1
    row = db_rows(db)[0]
    assert row["status"] == "エラー"
    assert "データがありません" in row["error"]
    assert row["rows"] is None


def test_folder_check_records_every_csv(tmp_path):
    folder = tmp_path / "in"
    folder.mkdir()
    write(folder / "ok.csv", "id\n1\n")
    write(folder / "empty.csv", "")
    db = tmp_path / "h.db"
    main(["--folder", str(folder), "-o", str(tmp_path / "out"), "--history", str(db)])
    rows = {r["file_name"]: r["status"] for r in db_rows(db)}
    assert rows == {"ok.csv": "問題なし", "empty.csv": "エラー"}


def test_without_history_option_no_database_is_created(tmp_path):
    csv = write(tmp_path / "a.csv", "id\n1\n")
    main([str(csv), "-o", str(tmp_path / "r.xlsx")])
    assert [p.name for p in tmp_path.iterdir() if p.suffix == ".db"] == []


def test_unusable_history_path_stops_before_checking(tmp_path, capsys):
    csv = write(tmp_path / "a.csv", "id\n1\n")
    output = tmp_path / "r.xlsx"
    code = main([str(csv), "-o", str(output), "--history", str(tmp_path / "nofolder" / "h.db")])
    assert code == 1
    assert "見つかりません" in capsys.readouterr().out
    assert not output.exists()


def test_history_path_that_is_not_a_database_is_an_error(tmp_path, capsys):
    csv = write(tmp_path / "a.csv", "id\n1\n")
    bad = write(tmp_path / "bad.db", "this is not sqlite" * 50)
    output = tmp_path / "r.xlsx"
    code = main([str(csv), "-o", str(output), "--history", str(bad)])
    assert code == 1
    assert "エラー" in capsys.readouterr().out
    assert not output.exists()


def test_show_history_prints_changes_from_previous_run(tmp_path, capsys):
    csv = write(tmp_path / "a.csv", "id,name\n1,A\n2,B\n")
    db = tmp_path / "h.db"
    main([str(csv), "-o", str(tmp_path / "r.xlsx"), "--history", str(db)])
    write(csv, "id,name\n1,A\n1,A\n2,\n")
    main([str(csv), "-o", str(tmp_path / "r.xlsx"), "--history", str(db)])
    capsys.readouterr()

    code = main(["--show-history", str(db)])
    out = capsys.readouterr().out
    assert code == 0
    lines = out.strip().splitlines()
    assert "[初回]" in lines[0]
    assert "判定 問題なし→問題あり" in lines[1]
    assert "データ件数 2→3(+1)" in lines[1]
    assert "欠損セル数 0→1(+1)" in lines[1]
    assert "重複件数 0→1(+1)" in lines[1]
    assert "合計 2 件の記録" in lines[-1]


def test_same_name_in_different_folders_is_not_compared(tmp_path):
    db = tmp_path / "h.db"
    for name, text in (("x", "id\n1\n"), ("y", "id\n1\n1\n")):
        folder = tmp_path / name
        folder.mkdir()
        csv = write(folder / "a.csv", text)
        main([str(csv), "-o", str(tmp_path / "r.xlsx"), "--history", str(db)])
    lines = format_history(load_history(db))
    assert all("[初回]" in line for line in lines)


def test_show_history_empty_database(tmp_path, capsys):
    db = tmp_path / "h.db"
    open_history(db).close()
    assert main(["--show-history", str(db)]) == 0
    assert "記録がありません" in capsys.readouterr().out


def test_show_history_missing_database_is_error_and_creates_nothing(tmp_path, capsys):
    db = tmp_path / "none.db"
    assert main(["--show-history", str(db)]) == 1
    assert "見つかりません" in capsys.readouterr().out
    assert not db.exists()


def test_show_history_cannot_be_combined_with_a_check(tmp_path, capsys):
    csv = write(tmp_path / "a.csv", "id\n1\n")
    db = tmp_path / "h.db"
    open_history(db).close()
    assert main([str(csv), "--show-history", str(db)]) == 1
    assert "同時に使えません" in capsys.readouterr().out


def test_describe_changes_same_and_error_cases():
    base = {"status": "問題なし", "rows": 2, "columns": 2, "missing": 0, "duplicates": 0, "violations": None}
    assert describe_changes(None, base) == "初回"
    assert describe_changes(base, dict(base)) == "前回と同じ"
    error = {**base, "status": "エラー", "rows": None, "columns": None, "missing": None, "duplicates": None}
    # 数字がない側(エラー)とは件数を比べず、判定の変化だけ示す
    assert describe_changes(base, error) == "判定 問題なし→エラー"


def test_load_history_can_filter_and_limit(tmp_path):
    db = tmp_path / "h.db"
    a = write(tmp_path / "a.csv", "id\n1\n")
    b = write(tmp_path / "b.csv", "id\n1\n")
    for csv in (a, a, a, b):
        main([str(csv), "-o", str(tmp_path / "r.xlsx"), "--history", str(db)])
    assert len(load_history(db)) == 4
    assert len(load_history(db, file_name="a.csv")) == 3
    limited = load_history(db, limit=2)
    assert sorted(r["file_name"] for r in limited) == ["a.csv", "a.csv", "b.csv"]


def test_load_history_missing_file_raises(tmp_path):
    with pytest.raises(HistoryError):
        load_history(tmp_path / "nothing.db")


def make_history(tmp_path):
    """a.csvを3回(欠損が増える)、b.csvを1回検査した記録を作る。"""
    db = tmp_path / "h.db"
    a = tmp_path / "a.csv"
    b = write(tmp_path / "b.csv", "id\n1\n")
    for text in ("id,name\n1,A\n2,B\n", "id,name\n1,A\n2,\n", "id,name\n1,\n2,\n"):
        write(a, text)
        main([str(a), "-o", str(tmp_path / "r.xlsx"), "--history", str(db)])
    main([str(b), "-o", str(tmp_path / "r.xlsx"), "--history", str(db)])
    return db


def test_history_name_shows_only_that_csv(tmp_path, capsys):
    db = make_history(tmp_path)
    capsys.readouterr()
    assert main(["--show-history", str(db), "--history-name", "b.csv"]) == 0
    out = capsys.readouterr().out
    assert "b.csv" in out and "a.csv" not in out
    assert "合計 1 件" in out


def test_history_name_without_match_says_so(tmp_path, capsys):
    db = make_history(tmp_path)
    capsys.readouterr()
    assert main(["--show-history", str(db), "--history-name", "zzz.csv"]) == 0
    assert "zzz.csv の記録がありません" in capsys.readouterr().out


def test_history_limit_keeps_newest_per_csv(tmp_path, capsys):
    db = make_history(tmp_path)
    capsys.readouterr()
    assert main(["--show-history", str(db), "--history-limit", "1"]) == 0
    lines = capsys.readouterr().out.strip().splitlines()
    assert "合計 2 件" in lines[-1]
    a_line = [line for line in lines if "a.csv" in line][0]
    assert "欠損セル数=2" in a_line  # 3回目(最新)が残る
    # 絞り込んでも、前回(2回目)との違いが出る(「初回」にはならない)
    assert "欠損セル数 1→2(+1)" in a_line
    assert "[初回]" not in a_line


def test_history_limit_must_be_positive(tmp_path, capsys):
    db = make_history(tmp_path)
    capsys.readouterr()
    assert main(["--show-history", str(db), "--history-limit", "0"]) == 1
    assert "1以上" in capsys.readouterr().out


def test_history_filter_options_need_show_history(tmp_path, capsys):
    csv = write(tmp_path / "a.csv", "id\n1\n")
    for extra in (["--history-name", "a.csv"], ["--history-limit", "2"], ["--history-excel", str(tmp_path / "x.xlsx")]):
        assert main([str(csv), "-o", str(tmp_path / "r.xlsx")] + extra) == 1
        assert "--show-history と一緒に使います" in capsys.readouterr().out


def test_history_excel_has_rows_and_trend_chart_for_one_csv(tmp_path):
    from openpyxl import load_workbook

    db = make_history(tmp_path)
    out = tmp_path / "history.xlsx"
    assert main(["--show-history", str(db), "--history-name", "a.csv", "--history-excel", str(out)]) == 0
    book = load_workbook(out)
    rows = [[c.value for c in row] for row in book["履歴"].iter_rows()]
    assert rows[0][:3] == ["検査日時", "ファイル名", "判定"]
    assert len(rows) == 4  # 見出し+3回分
    assert rows[1][8] == "初回"
    assert "欠損セル数 0→1(+1)" in rows[2][8]
    assert [r[5] for r in rows[1:]] == [0, 1, 2]
    trend = book["推移"]
    assert len(trend._charts) == 1


def test_history_excel_has_no_chart_for_several_csvs(tmp_path):
    from openpyxl import load_workbook

    db = make_history(tmp_path)
    out = tmp_path / "history.xlsx"
    assert main(["--show-history", str(db), "--history-excel", str(out)]) == 0
    book = load_workbook(out)
    assert book.sheetnames == ["履歴"]
    assert len(list(book["履歴"].iter_rows())) == 5


def test_history_excel_cannot_overwrite_the_database(tmp_path, capsys):
    db = make_history(tmp_path)
    capsys.readouterr()
    assert main(["--show-history", str(db), "--history-excel", str(db)]) == 1
    assert "同じ" in capsys.readouterr().out
    assert len(load_history(db)) == 4  # データベースは無事


def test_history_excel_path_may_exist_while_database_is_missing(tmp_path, capsys):
    out = write(tmp_path / "x.xlsx", "dummy")
    assert main(["--show-history", str(tmp_path / "none.db"), "--history-excel", str(out)]) == 1
    assert "見つかりません" in capsys.readouterr().out
