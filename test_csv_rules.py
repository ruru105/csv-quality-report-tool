import json

import pytest
from openpyxl import load_workbook

from csv_report import main
from csv_rules import RulesError, load_rules


def write_csv(tmp_path, text, name="data.csv"):
    path = tmp_path / name
    path.write_text(text, encoding="utf-8-sig")
    return path


def write_rules(tmp_path, rules, name="rules.json"):
    path = tmp_path / name
    path.write_text(json.dumps(rules, ensure_ascii=False), encoding="utf-8")
    return path


def run(tmp_path, csv_text, rules):
    csv_file = write_csv(tmp_path, csv_text)
    rules_file = write_rules(tmp_path, rules)
    output = tmp_path / "out.xlsx"
    code = main([str(csv_file), "-c", str(rules_file), "-o", str(output)])
    return code, output


def sheet_rows(output, sheet):
    return [[cell.value for cell in row] for row in load_workbook(output)[sheet].iter_rows()]


def summary(output):
    return dict(sheet_rows(output, "検査結果")[1:])


# ---- 設定ファイルの読み込み ----


def test_load_rules_reads_all_keys(tmp_path):
    path = write_rules(
        tmp_path,
        {"required_columns": ["id"], "not_null_columns": ["id", "name"], "unique_columns": ["id"]},
    )
    rules = load_rules(path)
    assert rules["required_columns"] == ["id"]
    assert rules["not_null_columns"] == ["id", "name"]
    assert rules["unique_columns"] == ["id"]


def test_load_rules_empty_object_means_default(tmp_path):
    rules = load_rules(write_rules(tmp_path, {}))
    assert rules == {"required_columns": [], "not_null_columns": None, "unique_columns": None}


def test_load_rules_missing_file(tmp_path):
    with pytest.raises(RulesError, match="見つかりません"):
        load_rules(tmp_path / "nothing.json")


def test_load_rules_invalid_json(tmp_path):
    path = tmp_path / "bad.json"
    path.write_text("{ required_columns: ", encoding="utf-8")
    with pytest.raises(RulesError, match="JSONとして読めません"):
        load_rules(path)


def test_load_rules_not_an_object(tmp_path):
    with pytest.raises(RulesError, match="{ ... } の形"):
        load_rules(write_rules(tmp_path, ["id"]))


def test_load_rules_unknown_key_is_rejected(tmp_path):
    # 綴りの間違いが、黙って無視されないこと。
    with pytest.raises(RulesError, match="未対応の項目.*unique_column"):
        load_rules(write_rules(tmp_path, {"unique_column": ["id"]}))


@pytest.mark.parametrize("value", [[], "id", [1], [""], None])
def test_load_rules_bad_value_is_rejected(tmp_path, value):
    with pytest.raises(RulesError, match="required_columns"):
        load_rules(write_rules(tmp_path, {"required_columns": value}))


# ---- 検査への反映 ----


def test_missing_required_column_is_a_problem(tmp_path):
    code, output = run(tmp_path, "id,name\n1,Aoki\n2,Sato\n", {"required_columns": ["id", "email"]})
    assert code == 0
    assert summary(output)["総合判定"] == "問題あり"
    rules = {row[0]: row for row in sheet_rows(output, "ルール判定")[1:]}
    assert rules["必須列"][2] == "NG"
    assert "email" in rules["必須列"][3]


def test_rules_satisfied_is_ok_even_if_other_columns_have_gaps(tmp_path):
    # 欠損セル数は全体の数字のまま表示するが、判定はルール(idとname)だけで行う。
    csv_text = "id,name,memo\n1,Aoki,\n2,Sato,\n"
    code, output = run(tmp_path, csv_text, {"not_null_columns": ["id", "name"]})
    assert code == 0
    result = summary(output)
    assert result["欠損セル数"] == 2
    assert result["ルール違反数"] == 0
    assert result["総合判定"] == "問題なし"


def test_not_null_column_with_gap_is_a_problem(tmp_path):
    code, output = run(tmp_path, "id,name\n1,Aoki\n,Sato\n", {"not_null_columns": ["id"]})
    assert summary(output)["総合判定"] == "問題あり"
    assert sheet_rows(output, "欠損行")[1:] == [[None, "Sato"]]
    rules = {row[0]: row for row in sheet_rows(output, "ルール判定")[1:]}
    assert "id=1" in rules["欠損を許さない列"][3]


def test_duplicate_is_judged_by_unique_columns_only(tmp_path):
    # 名前が違っても、idが同じなら重複。
    csv_text = "id,name\n1,Aoki\n1,Sato\n2,Tanaka\n"
    code, output = run(tmp_path, csv_text, {"unique_columns": ["id"]})
    assert summary(output)["総合判定"] == "問題あり"
    assert len(sheet_rows(output, "重複行")) == 3  # 見出し+2行
    # 行全体で見る既定では、重複なし(元の動き)。
    code, output = run(tmp_path, csv_text, {})
    assert summary(output)["総合判定"] == "問題なし"


def test_unique_columns_are_compared_as_text(tmp_path):
    code, output = run(tmp_path, "id,name\n007,A\n7,B\n", {"unique_columns": ["id"]})
    assert summary(output)["総合判定"] == "問題なし"


def test_rule_column_not_in_csv_is_reported_not_crashed(tmp_path):
    code, output = run(
        tmp_path,
        "id,name\n1,Aoki\n",
        {"not_null_columns": ["id", "phone"], "unique_columns": ["zip"]},
    )
    assert code == 0
    assert summary(output)["総合判定"] == "問題あり"
    details = " ".join(row[3] for row in sheet_rows(output, "ルール判定")[1:])
    assert "phone" in details and "zip" in details


def test_empty_rules_behave_like_default_judgement(tmp_path):
    code, output = run(tmp_path, "id,name\n1,Aoki\n1,Aoki\n2,\n", {})
    assert summary(output)["総合判定"] == "問題あり"
    sheets = load_workbook(output).sheetnames
    assert sheets[:2] == ["検査結果", "ルール判定"]


def test_without_config_there_is_no_rule_sheet(tmp_path):
    csv_file = write_csv(tmp_path, "id,name\n1,Aoki\n")
    output = tmp_path / "out.xlsx"
    assert main([str(csv_file), "-o", str(output)]) == 0
    wb = load_workbook(output)
    assert "ルール判定" not in wb.sheetnames
    assert "ルール違反数" not in summary(output)


def test_invalid_config_stops_without_creating_excel(tmp_path, capsys):
    csv_file = write_csv(tmp_path, "id\n1\n")
    bad = write_rules(tmp_path, {"unique_column": ["id"]})
    output = tmp_path / "out.xlsx"
    assert main([str(csv_file), "-c", str(bad), "-o", str(output)]) == 1
    assert not output.exists()
    assert "エラー" in capsys.readouterr().out


def test_missing_config_file_stops(tmp_path, capsys):
    csv_file = write_csv(tmp_path, "id\n1\n")
    output = tmp_path / "out.xlsx"
    assert main([str(csv_file), "-c", str(tmp_path / "none.json"), "-o", str(output)]) == 1
    assert not output.exists()
    assert "見つかりません" in capsys.readouterr().out
