import pandas as pd

from csv_clean import main


def write_csv(path, text, encoding="utf-8-sig"):
    path.write_text(text, encoding=encoding)
    return path


def read_output(path):
    """整形後のCSVを、すべて文字のまま読み込む。"""
    return pd.read_csv(path, encoding="utf-8-sig", dtype=str, keep_default_na=False)


def test_trim_spaces_and_remove_duplicates(tmp_path, capsys):
    # 前後の空白(全角を含む)を消したあと、同じ内容になった行が取り除かれること。
    input_file = write_csv(
        tmp_path / "input.csv",
        "id,name\n1, Aoki \n2,　Sato\n2,Sato\n1,Aoki\n",
    )
    output_file = tmp_path / "output.csv"

    result = main([str(input_file), "-o", str(output_file)])

    assert result == 0
    assert read_output(output_file).to_dict("records") == [
        {"id": "1", "name": "Aoki"},
        {"id": "2", "name": "Sato"},
    ]

    output = capsys.readouterr().out
    assert "データ件数=4→2" in output
    assert "空白を消した箇所=2" in output
    assert "取り除いた重複行=2" in output


def test_header_spaces_are_removed(tmp_path, capsys):
    input_file = write_csv(tmp_path / "input.csv", " id , name \n1,Aoki\n")
    output_file = tmp_path / "output.csv"

    main([str(input_file), "-o", str(output_file)])

    assert list(read_output(output_file).columns) == ["id", "name"]
    assert "空白を消した箇所=2" in capsys.readouterr().out


def test_original_file_is_not_changed(tmp_path):
    input_file = write_csv(tmp_path / "input.csv", "id,name\n1, Aoki \n1, Aoki \n")
    before = input_file.read_bytes()

    main([str(input_file), "-o", str(tmp_path / "output.csv")])

    assert input_file.read_bytes() == before


def test_values_are_kept_as_written(tmp_path):
    # 先頭の0、空欄、「NA」という文字が、勝手に変わらないこと。
    input_file = write_csv(
        tmp_path / "input.csv",
        "code,name,memo\n007,Aoki,\n008,NA,ok\n",
    )
    output_file = tmp_path / "output.csv"

    main([str(input_file), "-o", str(output_file)])

    assert read_output(output_file).to_dict("records") == [
        {"code": "007", "name": "Aoki", "memo": ""},
        {"code": "008", "name": "NA", "memo": "ok"},
    ]


def test_default_output_name(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    write_csv(tmp_path / "input.csv", "id,name\n1,Aoki\n")

    main(["input.csv"])

    assert (tmp_path / "cleaned.csv").exists()


def test_cp932_input_is_saved_as_utf8_with_bom(tmp_path):
    input_file = write_csv(
        tmp_path / "input.csv",
        "id,name\n1,青木\n1,青木\n",
        encoding="cp932",
    )
    output_file = tmp_path / "output.csv"

    main([str(input_file), "-o", str(output_file)])

    assert output_file.read_bytes().startswith(b"\xef\xbb\xbf")
    assert read_output(output_file).to_dict("records") == [
        {"id": "1", "name": "青木"},
    ]


def test_header_only_csv_is_saved_as_header_only(tmp_path):
    input_file = write_csv(tmp_path / "input.csv", "id,name\n")
    output_file = tmp_path / "output.csv"

    result = main([str(input_file), "-o", str(output_file)])

    assert result == 0
    saved = read_output(output_file)
    assert list(saved.columns) == ["id", "name"]
    assert len(saved) == 0


def test_missing_input_file_shows_error(tmp_path, capsys):
    output_file = tmp_path / "output.csv"

    result = main([str(tmp_path / "no_such_file.csv"), "-o", str(output_file)])

    assert result == 1
    assert "見つかりません" in capsys.readouterr().out
    assert not output_file.exists()


def test_same_input_and_output_is_refused(tmp_path, capsys):
    # 元のファイルを上書きしないこと。
    input_file = write_csv(tmp_path / "input.csv", "id,name\n1, Aoki \n")
    before = input_file.read_bytes()

    result = main([str(input_file), "-o", str(input_file)])

    assert result == 1
    assert "元のファイルを書き換えない" in capsys.readouterr().out
    assert input_file.read_bytes() == before


def test_empty_csv_shows_error(tmp_path, capsys):
    input_file = write_csv(tmp_path / "empty.csv", "")
    output_file = tmp_path / "output.csv"

    result = main([str(input_file), "-o", str(output_file)])

    assert result == 1
    assert "データがありません" in capsys.readouterr().out
    assert not output_file.exists()


def test_unwritable_output_shows_error(tmp_path, capsys):
    # 保存先のフォルダがないときに、分かりやすいエラーを出すこと。
    input_file = write_csv(tmp_path / "input.csv", "id,name\n1,Aoki\n")

    result = main([str(input_file), "-o", str(tmp_path / "no_folder" / "output.csv")])

    assert result == 1
    assert "保存できませんでした" in capsys.readouterr().out


def test_where_keeps_only_matching_rows(tmp_path, capsys):
    input_file = write_csv(
        tmp_path / "input.csv",
        "id,department\n1,Care\n2,Office\n3,Care\n",
    )
    output_file = tmp_path / "output.csv"

    result = main(
        [str(input_file), "-o", str(output_file), "--where", "department=Care"]
    )

    assert result == 0
    assert read_output(output_file).to_dict("records") == [
        {"id": "1", "department": "Care"},
        {"id": "3", "department": "Care"},
    ]
    assert "データ件数=3→2" in capsys.readouterr().out


def test_where_is_exact_match_not_partial(tmp_path):
    # 「Care」を指定したとき、「Care部」のような部分一致は含めないこと。
    input_file = write_csv(
        tmp_path / "input.csv",
        "id,department\n1,Care\n2,Care部\n",
    )
    output_file = tmp_path / "output.csv"

    main([str(input_file), "-o", str(output_file), "--where", "department=Care"])

    assert read_output(output_file).to_dict("records") == [
        {"id": "1", "department": "Care"},
    ]


def test_where_no_match_creates_header_only_csv(tmp_path, capsys):
    input_file = write_csv(tmp_path / "input.csv", "id,department\n1,Care\n")
    output_file = tmp_path / "output.csv"

    result = main(
        [str(input_file), "-o", str(output_file), "--where", "department=Office"]
    )

    assert result == 0
    saved = read_output(output_file)
    assert list(saved.columns) == ["id", "department"]
    assert len(saved) == 0
    assert "データ件数=1→0" in capsys.readouterr().out


def test_where_invalid_format_shows_error(tmp_path, capsys):
    input_file = write_csv(tmp_path / "input.csv", "id,department\n1,Care\n")
    output_file = tmp_path / "output.csv"

    result = main(
        [str(input_file), "-o", str(output_file), "--where", "department:Care"]
    )

    assert result == 1
    assert "正しくありません" in capsys.readouterr().out
    assert not output_file.exists()


def test_where_unknown_column_shows_error(tmp_path, capsys):
    input_file = write_csv(tmp_path / "input.csv", "id,department\n1,Care\n")
    output_file = tmp_path / "output.csv"

    result = main(
        [str(input_file), "-o", str(output_file), "--where", "no_such_column=Care"]
    )

    assert result == 1
    assert "見つかりません" in capsys.readouterr().out
    assert not output_file.exists()


def test_drop_single_column(tmp_path, capsys):
    input_file = write_csv(
        tmp_path / "input.csv",
        "id,name,memo\n1,Aoki,内緒\n",
    )
    output_file = tmp_path / "output.csv"

    result = main([str(input_file), "-o", str(output_file), "--drop", "memo"])

    assert result == 0
    assert list(read_output(output_file).columns) == ["id", "name"]
    assert "削除した列=memo" in capsys.readouterr().out


def test_drop_multiple_columns_with_comma(tmp_path):
    input_file = write_csv(
        tmp_path / "input.csv",
        "id,name,memo,note\n1,Aoki,内緒,秘密\n",
    )
    output_file = tmp_path / "output.csv"

    main([str(input_file), "-o", str(output_file), "--drop", "memo,note"])

    assert list(read_output(output_file).columns) == ["id", "name"]


def test_drop_unknown_column_shows_error(tmp_path, capsys):
    input_file = write_csv(tmp_path / "input.csv", "id,name\n1,Aoki\n")
    output_file = tmp_path / "output.csv"

    result = main([str(input_file), "-o", str(output_file), "--drop", "no_such_column"])

    assert result == 1
    assert "見つかりません" in capsys.readouterr().out
    assert not output_file.exists()


def test_drop_all_columns_is_refused(tmp_path, capsys):
    input_file = write_csv(tmp_path / "input.csv", "id,name\n1,Aoki\n")
    output_file = tmp_path / "output.csv"

    result = main([str(input_file), "-o", str(output_file), "--drop", "id,name"])

    assert result == 1
    assert "すべての列を削除することはできません" in capsys.readouterr().out
    assert not output_file.exists()


def test_where_and_drop_together(tmp_path):
    # --where は絞り込みに使った列を、--drop であとから削除できること。
    input_file = write_csv(
        tmp_path / "input.csv",
        "id,department,memo\n1,Care,内緒\n2,Office,内緒\n",
    )
    output_file = tmp_path / "output.csv"

    main(
        [
            str(input_file),
            "-o",
            str(output_file),
            "--where",
            "department=Care",
            "--drop",
            "department,memo",
        ]
    )

    assert read_output(output_file).to_dict("records") == [{"id": "1"}]


def test_tab_separated_input_is_split_into_columns(tmp_path):
    # FX取引ソフトの出力など、タブ区切りのファイルも正しく列を分けて整形できること。
    input_file = write_csv(tmp_path / "input.csv", "id\tname\n1\tAoki\n2\tSato\n")
    output_file = tmp_path / "output.csv"

    result = main([str(input_file), "-o", str(output_file)])

    assert result == 0
    assert read_output(output_file).to_dict("records") == [
        {"id": "1", "name": "Aoki"},
        {"id": "2", "name": "Sato"},
    ]


def test_semicolon_separated_input_is_split_into_columns(tmp_path):
    input_file = write_csv(tmp_path / "input.csv", "id;name\n1;Aoki\n2;Sato\n")
    output_file = tmp_path / "output.csv"

    result = main([str(input_file), "-o", str(output_file)])

    assert result == 0
    assert read_output(output_file).to_dict("records") == [
        {"id": "1", "name": "Aoki"},
        {"id": "2", "name": "Sato"},
    ]


def test_genuinely_single_column_input_stays_single_column(tmp_path):
    # 区切り文字を含まない、もともと1列だけのCSVを誤って分割しないこと。
    input_file = write_csv(tmp_path / "input.csv", "name\nAoki\nSato\n")
    output_file = tmp_path / "output.csv"

    result = main([str(input_file), "-o", str(output_file)])

    assert result == 0
    assert read_output(output_file).to_dict("records") == [
        {"name": "Aoki"},
        {"name": "Sato"},
    ]


# ---- 並べ替え(--sort・--desc)と列名の変更(--rename) ----

def run_clean(tmp_path, text, *options):
    input_file = write_csv(tmp_path / "input.csv", text)
    output_file = tmp_path / "output.csv"
    result = main([str(input_file), "-o", str(output_file), *options])
    return result, output_file


def test_sort_numeric_column_by_size_not_by_text(tmp_path):
    # 文字の順だと「100」が「30」より先に来てしまうので、数の大きさで並べる。
    result, out = run_clean(tmp_path, "id,amount\n1,100\n2,30\n3,50\n", "--sort", "amount")
    assert result == 0
    assert list(read_output(out)["amount"]) == ["30", "50", "100"]


def test_sort_descending(tmp_path):
    result, out = run_clean(tmp_path, "id,amount\n1,100\n2,30\n3,50\n", "--sort", "amount", "--desc")
    assert result == 0
    assert list(read_output(out)["amount"]) == ["100", "50", "30"]


def test_sort_keeps_leading_zeros_in_values(tmp_path):
    # 並べ替えでも、「007」は「007」のまま。
    result, out = run_clean(tmp_path, "id,n\n007,2\n003,1\n", "--sort", "n")
    assert result == 0
    assert list(read_output(out)["id"]) == ["003", "007"]


def test_sort_text_column_uses_text_order(tmp_path):
    result, out = run_clean(tmp_path, "id,name\n1,Sato\n2,Aoki\n3,Tanaka\n", "--sort", "name")
    assert result == 0
    assert list(read_output(out)["name"]) == ["Aoki", "Sato", "Tanaka"]


def test_sort_puts_blank_values_last_in_both_directions(tmp_path):
    text = "id,amount\n1,\n2,30\n3,100\n"
    _, up = run_clean(tmp_path, text, "--sort", "amount")
    assert list(read_output(up)["id"]) == ["2", "3", "1"]
    _, down = run_clean(tmp_path, text, "--sort", "amount", "--desc")
    assert list(read_output(down)["id"]) == ["3", "2", "1"]


def test_sort_keeps_original_order_for_equal_values(tmp_path):
    text = "id,group\nA,2\nB,1\nC,2\nD,1\n"
    _, up = run_clean(tmp_path, text, "--sort", "group")
    assert list(read_output(up)["id"]) == ["B", "D", "A", "C"]
    _, down = run_clean(tmp_path, text, "--sort", "group", "--desc")
    assert list(read_output(down)["id"]) == ["A", "C", "B", "D"]


def test_sort_mixed_numbers_and_text_falls_back_to_text_order(tmp_path):
    result, out = run_clean(tmp_path, "id,v\n1,10\n2,9\n3,abc\n", "--sort", "v")
    assert result == 0
    assert list(read_output(out)["v"]) == ["10", "9", "abc"]


def test_sort_unknown_column_is_an_error_and_no_file(tmp_path, capsys):
    result, out = run_clean(tmp_path, "id,amount\n1,100\n", "--sort", "price")
    assert result == 1
    assert not out.exists()
    assert "並べ替えに使う列名「price」がCSVに見つかりません" in capsys.readouterr().out


def test_desc_without_sort_is_an_error(tmp_path, capsys):
    result, out = run_clean(tmp_path, "id,amount\n1,100\n", "--desc")
    assert result == 1
    assert not out.exists()
    assert "--desc は --sort と一緒に" in capsys.readouterr().out


def test_rename_changes_only_the_header(tmp_path, capsys):
    result, out = run_clean(tmp_path, "id,name,dept\n007,Aoki,Care\n", "--rename", "name=氏名,dept=所属")
    assert result == 0
    data = read_output(out)
    assert list(data.columns) == ["id", "氏名", "所属"]
    assert data.to_dict("records") == [{"id": "007", "氏名": "Aoki", "所属": "Care"}]
    assert "列名の変更=name→氏名、dept→所属" in capsys.readouterr().out


def test_rename_unknown_column_is_an_error_and_no_file(tmp_path, capsys):
    result, out = run_clean(tmp_path, "id,name\n1,Aoki\n", "--rename", "title=氏名")
    assert result == 1
    assert not out.exists()
    assert "名前を変える列がCSVに見つかりません：title" in capsys.readouterr().out


def test_rename_to_an_existing_column_name_is_an_error(tmp_path, capsys):
    result, out = run_clean(tmp_path, "id,name\n1,Aoki\n", "--rename", "name=id")
    assert result == 1
    assert not out.exists()
    assert "ほかの列と重なります" in capsys.readouterr().out


def test_rename_two_columns_to_the_same_name_is_an_error(tmp_path):
    result, out = run_clean(tmp_path, "a,b\n1,2\n", "--rename", "a=x,b=x")
    assert result == 1
    assert not out.exists()


def test_rename_same_column_twice_is_an_error(tmp_path):
    result, out = run_clean(tmp_path, "a,b\n1,2\n", "--rename", "a=x,a=y")
    assert result == 1
    assert not out.exists()


def test_rename_swapping_two_names_is_allowed(tmp_path):
    # 入れ替えは、重なりにならない。
    result, out = run_clean(tmp_path, "a,b\n1,2\n", "--rename", "a=b,b=a")
    assert result == 0
    assert read_output(out).to_dict("records") == [{"b": "1", "a": "2"}]


def test_rename_bad_format_is_an_error(tmp_path, capsys):
    for bad in ["name", "name=", "=氏名", ","]:
        result, out = run_clean(tmp_path, "id,name\n1,Aoki\n", "--rename", bad)
        assert result == 1
        assert not out.exists()
    assert "--rename の指定" in capsys.readouterr().out


def test_where_drop_sort_use_original_names_and_rename_is_applied_last(tmp_path):
    text = "id,dept,amount\n1,Care,50\n2,Office,10\n3,Care,100\n"
    result, out = run_clean(
        tmp_path, text,
        "--where", "dept=Care", "--sort", "amount", "--desc",
        "--drop", "dept", "--rename", "amount=金額",
    )
    assert result == 0
    assert read_output(out).to_dict("records") == [
        {"id": "3", "金額": "100"},
        {"id": "1", "金額": "50"},
    ]


def test_rename_a_dropped_column_is_an_error(tmp_path):
    result, out = run_clean(tmp_path, "id,dept\n1,Care\n", "--drop", "dept", "--rename", "dept=所属")
    assert result == 1
    assert not out.exists()


def test_sort_and_rename_do_not_change_the_input_file(tmp_path):
    text = "id,amount\n1,100\n2,30\n"
    input_file = write_csv(tmp_path / "input.csv", text)
    before = input_file.read_bytes()
    main([str(input_file), "-o", str(tmp_path / "o.csv"), "--sort", "amount", "--rename", "amount=金額"])
    assert input_file.read_bytes() == before


# ---- Excel(.xlsx)の読み込みと書き出し ----

import datetime

import pytest
from openpyxl import Workbook, load_workbook


def make_xlsx(path, rows, sheet="売上", extra_sheets=()):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for row in rows:
        worksheet.append(row)
    for name, extra_rows in extra_sheets:
        extra = workbook.create_sheet(name)
        for row in extra_rows:
            extra.append(row)
    workbook.save(path)
    return path


def run_xlsx(tmp_path, rows, out_name="output.csv", *options, **kwargs):
    input_file = make_xlsx(tmp_path / "input.xlsx", rows, **kwargs)
    output_file = tmp_path / out_name
    result = main([str(input_file), "-o", str(output_file), *options])
    return result, output_file


def test_excel_input_is_read_as_text_and_saved_as_csv(tmp_path):
    rows = [
        ["id", "name", "amount", "day", "flag"],
        ["007", " Aoki ", 100, datetime.date(2026, 10, 1), True],
        ["003", "Sato", 50.5, datetime.datetime(2026, 10, 2, 9, 30), False],
        ["001", "Ito", 30.0, None, None],
    ]
    result, out = run_xlsx(tmp_path, rows)
    assert result == 0
    assert read_output(out).to_dict("records") == [
        {"id": "007", "name": "Aoki", "amount": "100", "day": "2026-10-01", "flag": "TRUE"},
        {"id": "003", "name": "Sato", "amount": "50.5", "day": "2026-10-02 09:30:00", "flag": "FALSE"},
        {"id": "001", "name": "Ito", "amount": "30", "day": "", "flag": ""},
    ]


def test_excel_input_uses_first_sheet_by_default_and_sheet_option_selects_another(tmp_path):
    rows = [["a"], ["first"]]
    extra = [("2枚目", [["a"], ["second"]])]
    _, first = run_xlsx(tmp_path, rows, "o1.csv", extra_sheets=extra)
    assert list(read_output(first)["a"]) == ["first"]
    result, second = run_xlsx(tmp_path, rows, "o2.csv", "--sheet", "2枚目", extra_sheets=extra)
    assert result == 0
    assert list(read_output(second)["a"]) == ["second"]


def test_excel_unknown_sheet_is_an_error_listing_sheets(tmp_path, capsys):
    result, out = run_xlsx(tmp_path, [["a"], ["1"]], "o.csv", "--sheet", "nothing")
    assert result == 1
    assert not out.exists()
    message = capsys.readouterr().out
    assert "シート「nothing」が見つかりません" in message
    assert "売上" in message


def test_sheet_option_with_csv_input_is_an_error(tmp_path, capsys):
    result, out = run_clean(tmp_path, "a\n1\n", "--sheet", "売上")
    assert result == 1
    assert not out.exists()
    assert "--sheet はExcelファイルを読むときだけ" in capsys.readouterr().out


def test_excel_blank_or_duplicate_headers_are_errors(tmp_path, capsys):
    result, out = run_xlsx(tmp_path, [["a", None, "c"], [1, 2, 3]])
    assert result == 1 and not out.exists()
    assert "見出し(1行目)が空の列があります" in capsys.readouterr().out
    result, out = run_xlsx(tmp_path, [["a", "a"], [1, 2]], "o2.csv")
    assert result == 1 and not out.exists()
    assert "同じ名前の見出しがあります" in capsys.readouterr().out


def test_excel_with_no_data_is_an_error(tmp_path, capsys):
    result, out = run_xlsx(tmp_path, [])
    assert result == 1 and not out.exists()
    assert "データがありません" in capsys.readouterr().out


def test_excel_header_only_gives_header_only_output(tmp_path):
    result, out = run_xlsx(tmp_path, [["a", "b"]])
    assert result == 0
    assert list(read_output(out).columns) == ["a", "b"]
    assert len(read_output(out)) == 0


def test_excel_trailing_empty_rows_and_columns_are_ignored(tmp_path):
    rows = [["a", "b", None], [1, 2, None], [None, None, None], [None, None, None]]
    result, out = run_xlsx(tmp_path, rows)
    assert result == 0
    assert read_output(out).to_dict("records") == [{"a": "1", "b": "2"}]


def test_broken_file_named_xlsx_is_an_error_not_a_crash(tmp_path, capsys):
    broken = tmp_path / "broken.xlsx"
    broken.write_text("これはExcelではありません", encoding="utf-8")
    out = tmp_path / "o.csv"
    assert main([str(broken), "-o", str(out)]) == 1
    assert not out.exists()
    assert "Excelファイルとして開けませんでした" in capsys.readouterr().out


def test_old_xls_format_is_an_error_with_a_hint(tmp_path, capsys):
    old = tmp_path / "old.xls"
    old.write_bytes(b"dummy")
    out = tmp_path / "o.csv"
    assert main([str(old), "-o", str(out)]) == 1
    assert not out.exists()
    assert ".xlsx」として保存し直してください" in capsys.readouterr().out


def test_formula_without_saved_result_reads_as_blank(tmp_path):
    # 数式のセルは、Excelが保存した計算結果を読む。計算結果がない(プログラムで作っただけの)ファイルは空欄。
    rows = [["a", "b"], [1, "=A2*2"]]
    result, out = run_xlsx(tmp_path, rows)
    assert result == 0
    assert read_output(out).to_dict("records") == [{"a": "1", "b": ""}]


def test_excel_input_file_is_not_changed(tmp_path):
    input_file = make_xlsx(tmp_path / "input.xlsx", [["a"], [1]])
    before = input_file.read_bytes()
    main([str(input_file), "-o", str(tmp_path / "o.xlsx"), "--rename", "a=b"])
    assert input_file.read_bytes() == before


def test_excel_input_and_output_same_path_is_an_error(tmp_path, capsys):
    input_file = make_xlsx(tmp_path / "input.xlsx", [["a"], [1]])
    assert main([str(input_file), "-o", str(input_file)]) == 1
    assert "保存先が、入力ファイル(CSVまたはExcel)と同じです" in capsys.readouterr().out


def test_csv_to_excel_creates_a_real_xlsx_with_header_style(tmp_path):
    result, out = run_clean(tmp_path, "id,name,amount\n007,Aoki,100\n003,Sato,50.5\n")
    assert result == 0
    out = tmp_path / "result.xlsx"
    input_file = tmp_path / "input.csv"
    assert main([str(input_file), "-o", str(out)]) == 0
    sheet = load_workbook(out).active
    assert out.read_bytes()[:2] == b"PK"            # 本物のxlsx(zip)。CSVの文字を別名で保存していない
    assert sheet.title == "整形結果"
    assert [c.value for c in sheet[1]] == ["id", "name", "amount"]
    assert sheet["A1"].font.bold is True
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:C3"
    # 普通の数は数として、先頭が0の値は文字のまま。
    assert sheet["A2"].value == "007" and sheet["A2"].data_type == "s"
    assert sheet["C2"].value == 100 and sheet["C2"].data_type == "n"
    assert sheet["C3"].value == 50.5


@pytest.mark.parametrize(
    "text",
    ["007", "1.0", "1.50", "+5", "1e3", "090-1234-5678", "12345678901234567", "-0", "0.", ".5", "1,000"],
)
def test_values_that_would_change_stay_text_in_excel_output(tmp_path, text):
    input_file = write_csv(tmp_path / "input.csv", f'v\n"{text}"\n')
    out = tmp_path / "out.xlsx"
    assert main([str(input_file), "-o", str(out)]) == 0
    cell = load_workbook(out).active["A2"]
    assert cell.value == text and cell.data_type == "s"


@pytest.mark.parametrize("text,expected", [("0", 0), ("12", 12), ("-3", -3), ("0.5", 0.5), ("100.25", 100.25)])
def test_plain_numbers_become_numbers_in_excel_output(tmp_path, text, expected):
    input_file = write_csv(tmp_path / "input.csv", f"v\n{text}\n")
    out = tmp_path / "out.xlsx"
    assert main([str(input_file), "-o", str(out)]) == 0
    cell = load_workbook(out).active["A2"]
    assert cell.value == expected and cell.data_type == "n"


def test_text_starting_with_equals_is_not_saved_as_a_formula(tmp_path):
    input_file = write_csv(tmp_path / "input.csv", 'v\n"=1+1"\n')
    out = tmp_path / "out.xlsx"
    assert main([str(input_file), "-o", str(out)]) == 0
    cell = load_workbook(out).active["A2"]
    assert cell.value == "=1+1" and cell.data_type == "s"


def test_excel_round_trip_keeps_every_value(tmp_path):
    values = ["007", "1.0", "", "100", "50.5", "-3", "=A1", "あ い", "090-1234"]
    csv_text = "v\n" + "\n".join(f'"{v}"' if v else '""' for v in values) + "\n"
    first = tmp_path / "first.xlsx"
    assert main([str(write_csv(tmp_path / "in.csv", csv_text, "utf-8")), "-o", str(first)]) == 0
    back = tmp_path / "back.csv"
    assert main([str(first), "-o", str(back)]) == 0
    assert list(read_output(back)["v"]) == values


def test_excel_to_excel_with_all_options(tmp_path):
    rows = [["id", "dept", "amount"], ["001", "Care", 50], ["002", "Office", 10], ["003", "Care", 100]]
    result, out = run_xlsx(
        tmp_path, rows, "out.xlsx",
        "--where", "dept=Care", "--sort", "amount", "--desc", "--drop", "dept", "--rename", "amount=売上額",
    )
    assert result == 0
    sheet = load_workbook(out).active
    assert [[c.value for c in row] for row in sheet.iter_rows()] == [
        ["id", "売上額"], ["003", 100], ["001", 50],
    ]


@pytest.mark.parametrize("name", ["out.xlsm", "out.xls"])
def test_unsupported_excel_output_formats_are_errors_and_no_file(tmp_path, name, capsys):
    input_file = write_csv(tmp_path / "input.csv", "a\n1\n")
    out = tmp_path / name
    assert main([str(input_file), "-o", str(out)]) == 1
    assert not out.exists()
    assert "には対応していません" in capsys.readouterr().out


def test_control_characters_cannot_be_written_to_excel_and_give_an_error(tmp_path, capsys):
    input_file = write_csv(tmp_path / "input.csv", "a\nx\x01y\n", "utf-8")
    out = tmp_path / "out.xlsx"
    assert main([str(input_file), "-o", str(out)]) == 1
    assert not out.exists()
    assert "制御文字" in capsys.readouterr().out


# ---- 条件で取り除く(--exclude)・外れた行数の表示・数式の注意 ----

ORDERS = "id,状態,金額\n1,確定,100\n2,キャンセル,200\n3,確定,300\n4,キャンセル待ち,400\n"


def test_exclude_removes_matching_rows_and_keeps_the_rest_in_order(tmp_path, capsys):
    result, out = run_clean(tmp_path, ORDERS, "--exclude", "状態=キャンセル")
    assert result == 0
    assert [r["id"] for r in read_output(out).to_dict("records")] == ["1", "3", "4"]
    assert "取り除いた行(--exclude)=1" in capsys.readouterr().out


def test_exclude_is_exact_match_not_partial(tmp_path):
    # 「キャンセル」を指定しても、「キャンセル待ち」は取り除かない。
    result, out = run_clean(tmp_path, ORDERS, "--exclude", "状態=キャンセル")
    assert "キャンセル待ち" in read_output(out)["状態"].tolist()


def test_exclude_with_no_match_keeps_every_row_and_reports_zero(tmp_path, capsys):
    result, out = run_clean(tmp_path, ORDERS, "--exclude", "状態=返品")
    assert result == 0
    assert len(read_output(out)) == 4
    assert "取り除いた行(--exclude)=0" in capsys.readouterr().out


def test_exclude_can_remove_every_row_and_leaves_a_header_only_file(tmp_path):
    result, out = run_clean(tmp_path, "id,状態\n1,キャンセル\n2,キャンセル\n", "--exclude", "状態=キャンセル")
    assert result == 0
    data = read_output(out)
    assert len(data) == 0
    assert list(data.columns) == ["id", "状態"]


def test_exclude_with_empty_value_removes_blank_cells_only(tmp_path):
    result, out = run_clean(tmp_path, "id,状態\n1,確定\n2,\n3,確定\n", "--exclude", "状態=")
    assert result == 0
    assert [r["id"] for r in read_output(out).to_dict("records")] == ["1", "3"]


def test_exclude_unknown_column_is_an_error_and_no_file(tmp_path, capsys):
    result, out = run_clean(tmp_path, ORDERS, "--exclude", "ステータス=キャンセル")
    assert result == 1
    assert not out.exists()
    assert "列名「ステータス」がCSVに見つかりません" in capsys.readouterr().out


def test_exclude_invalid_format_names_the_exclude_option(tmp_path, capsys):
    result, out = run_clean(tmp_path, ORDERS, "--exclude", "状態キャンセル")
    assert result == 1
    assert not out.exists()
    message = capsys.readouterr().out
    assert "--exclude の指定「状態キャンセル」が正しくありません" in message
    assert "--where" not in message


def test_where_and_exclude_together_and_both_counts_are_shown(tmp_path, capsys):
    text = "id,部署,状態\n1,A,確定\n2,A,キャンセル\n3,B,確定\n4,A,確定\n"
    result, out = run_clean(tmp_path, text, "--where", "部署=A", "--exclude", "状態=キャンセル")
    assert result == 0
    assert [r["id"] for r in read_output(out).to_dict("records")] == ["1", "4"]
    message = capsys.readouterr().out
    assert "絞り込み(--where)で外れた行=1" in message
    assert "取り除いた行(--exclude)=1" in message


def test_where_alone_reports_how_many_rows_it_removed(tmp_path, capsys):
    result, out = run_clean(tmp_path, ORDERS, "--where", "状態=確定")
    assert result == 0
    message = capsys.readouterr().out
    assert "絞り込み(--where)で外れた行=2" in message
    assert "--exclude" not in message


def test_counts_are_not_shown_when_options_are_not_used(tmp_path, capsys):
    run_clean(tmp_path, ORDERS)
    message = capsys.readouterr().out
    assert "--where" not in message and "--exclude" not in message


def test_exclude_uses_original_column_names_even_with_rename(tmp_path):
    result, out = run_clean(
        tmp_path, ORDERS, "--exclude", "状態=キャンセル", "--rename", "状態=Status"
    )
    assert result == 0
    data = read_output(out)
    assert list(data.columns) == ["id", "Status", "金額"]
    assert len(data) == 3


def test_exclude_runs_after_duplicate_removal_and_counts_stay_consistent(tmp_path, capsys):
    text = "id,状態\n1,確定\n1,確定\n2,キャンセル\n"
    result, out = run_clean(tmp_path, text, "--exclude", "状態=キャンセル")
    assert result == 0
    message = capsys.readouterr().out
    assert "データ件数=3→1" in message
    assert "取り除いた重複行=1" in message
    assert "取り除いた行(--exclude)=1" in message


def test_exclude_also_works_for_excel_input_and_output(tmp_path):
    rows = [["id", "状態"], [1, "確定"], [2, "キャンセル"], [3, "確定"]]
    input_file = make_xlsx(tmp_path / "input.xlsx", rows)
    out = tmp_path / "out.xlsx"
    assert main([str(input_file), "-o", str(out), "--exclude", "状態=キャンセル"]) == 0
    sheet = load_workbook(out).active
    assert [[c.value for c in row] for row in sheet.iter_rows()][1:] == [[1, "確定"], [3, "確定"]]


def test_same_path_message_names_both_csv_and_excel(tmp_path, capsys):
    input_file = write_csv(tmp_path / "input.csv", "a\n1\n")
    assert main([str(input_file), "-o", str(input_file)]) == 1
    assert "入力ファイル(CSVまたはExcel)と同じです" in capsys.readouterr().out


def test_formulas_without_saved_result_are_counted_in_a_notice(tmp_path, capsys):
    rows = [["a", "b"], [1, "=A2*2"], [2, "=A3*2"]]
    result, out = run_xlsx(tmp_path, rows)
    assert result == 0
    assert "計算結果が保存されていないセルが2個" in capsys.readouterr().out


def test_no_formula_notice_for_plain_excel_input(tmp_path, capsys):
    result, out = run_xlsx(tmp_path, [["a", "b"], [1, 2]])
    assert result == 0
    assert "計算結果" not in capsys.readouterr().out


def test_no_formula_notice_for_csv_input(tmp_path, capsys):
    run_clean(tmp_path, "a,b\n1,=A2*2\n")
    assert "計算結果" not in capsys.readouterr().out


def test_text_that_only_starts_with_equals_is_not_counted_as_a_formula(tmp_path, capsys):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["a"])
    cell = sheet.cell(row=2, column=1, value="=memo")
    cell.data_type = "s"
    path = tmp_path / "input.xlsx"
    workbook.save(path)
    out = tmp_path / "o.csv"
    assert main([str(path), "-o", str(out)]) == 0
    assert "計算結果" not in capsys.readouterr().out
    assert read_output(out).to_dict("records") == [{"a": "=memo"}]
